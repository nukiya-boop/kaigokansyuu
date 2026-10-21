"""js/data.js から確認テストの解答Excelを作る: python3 tools/make_answers.py js/data.js docs/確認テスト_解答一覧.xlsx"""
import sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.page import PageMargins

src, out = sys.argv[1:3]
t = open(src, encoding='utf-8').read()
D = json.loads(t[t.index('{'):t.rindex(';')])
months = D['months']

thin = Side(style='thin', color='999999')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill('solid', fgColor='DDEBE6')
legal_fill = PatternFill('solid', fgColor='FDF3DC')
wrap = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
bold = Font(bold=True)

def setup(ws, landscape=False, title_rows=None):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)
    if title_rows: ws.print_title_rows = title_rows

def header(ws, row, values):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = bold; c.fill = head_fill; c.border = border; c.alignment = center

wb = Workbook()

# 1. 正解一覧（採点用の早見表）
ws = wb.active; ws.title = '正解一覧'
ws['A1'] = D['fiscalYear'] + '　確認テスト 正解一覧（採点用）'; ws['A1'].font = Font(bold=True, size=14)
ws['A2'] = '問1〜5は○×、問6〜10はア・イ・ウ。1問1点・10点満点。7点以上が目安です。★は法定研修。'
header(ws, 4, ['月', '研修テーマ'] + ['問%d' % i for i in range(1, 11)])
for r, m in enumerate(months, 5):
    ws.cell(row=r, column=1, value=m['month']).alignment = center
    ws.cell(row=r, column=2, value=('★' if m['legal'] else '') + m['theme']).alignment = Alignment(vertical='center', wrap_text=True)
    for q in m['test']:
        c = ws.cell(row=r, column=2 + q['no'], value=q['answer'])
        c.alignment = center; c.font = Font(bold=True, size=13)
    for col in range(1, 13):
        ws.cell(row=r, column=col).border = border
        if m['legal']: ws.cell(row=r, column=col).fill = legal_fill
    ws.row_dimensions[r].height = 30
ws.column_dimensions['A'].width = 6
ws.column_dimensions['B'].width = 40
for i in range(10): ws.column_dimensions[chr(ord('C') + i)].width = 6
ws.freeze_panes = 'C5'
setup(ws, landscape=True)

# 2. 解答・解説（問題文・選択肢つき）
ws = wb.create_sheet('解答・解説')
ws['A1'] = '確認テスト 解答・解説（管理者・講師用）'; ws['A1'].font = Font(bold=True, size=14)
header(ws, 3, ['月', '問', '形式', '問題', '選択肢', '正解', '解説'])
r = 4
for m in months:
    ws.cell(row=r, column=1, value=m['month'] + '　' + ('★' if m['legal'] else '') + m['theme']).font = Font(bold=True, size=12)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    r += 1
    for q in m['test']:
        vals = [m['month'], q['no'], '○×' if q['type'] == 'ox' else '三択', q['q'],
                '\n'.join(c['k'] + '　' + c['t'] for c in q['choices']), q['answer'], q['explain']]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v); c.border = border; c.alignment = wrap
        for i in (1, 2, 3, 6): ws.cell(row=r, column=i).alignment = center
        ws.cell(row=r, column=6).font = Font(bold=True, size=13)
        r += 1
    r += 1
for col, w in zip('ABCDEFG', [6, 5, 6, 42, 30, 7, 46]): ws.column_dimensions[col].width = w
ws.freeze_panes = 'A4'
setup(ws, landscape=True, title_rows='3:3')

# 3. 研修資料の振り返りクイズ（各月5問）の解答
ws = wb.create_sheet('振り返りクイズ解答')
ws['A1'] = '研修資料 末尾の振り返りクイズ（○×5問）の解答'; ws['A1'].font = Font(bold=True, size=14)
header(ws, 3, ['月', '番号', '問題', '正解', '補足'])
r = 4
for m in months:
    for i, q in enumerate(m['review'], 1):
        vals = [m['month'], i, q['q'], q['a'], q['note']]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v); c.border = border; c.alignment = wrap
        for j in (1, 2, 4): ws.cell(row=r, column=j).alignment = center
        ws.cell(row=r, column=4).font = Font(bold=True, size=13)
        r += 1
for col, w in zip('ABCDE', [6, 6, 60, 7, 40]): ws.column_dimensions[col].width = w
ws.freeze_panes = 'A4'
setup(ws, title_rows='3:3')

# 4. 採点表（職員の解答を入れると自動で点数が出る）
ws = wb.create_sheet('採点表')
ws['A1'] = '採点表：月を選び、職員の解答（○×・アイウ）を入れると点数が自動計算されます'; ws['A1'].font = Font(bold=True, size=12)
ws['A2'] = '月'; ws['B2'] = months[0]['month']; ws['B2'].fill = PatternFill('solid', fgColor='FFFBE6'); ws['A2'].font = bold
from openpyxl.worksheet.datavalidation import DataValidation
dv = DataValidation(type='list', formula1='"' + ','.join(m['month'] for m in months) + '"', allow_blank=False)
ws.add_data_validation(dv); dv.add('B2')
header(ws, 4, ['氏名'] + ['問%d' % i for i in range(1, 11)] + ['点数', '判定'])
ws.cell(row=3, column=1, value='正解').font = bold
for i in range(10):
    col = 2 + i
    c = ws.cell(row=3, column=col, value='=INDEX(正解一覧!%s$5:%s$16,MATCH($B$2,正解一覧!$A$5:$A$16,0))' % ((chr(ord('C') + i),) * 2))
    c.alignment = center; c.font = Font(bold=True, color='2F7D6D'); c.border = border
input_fill = PatternFill('solid', fgColor='FFFBE6')
for r in range(5, 35):
    for col in range(1, 12):
        c = ws.cell(row=r, column=col); c.border = border; c.alignment = center
        if col <= 11: c.fill = input_fill
    ws.cell(row=r, column=12, value='=IF(COUNTA(B%d:K%d)=0,"",SUMPRODUCT(--(B%d:K%d=$B$3:$K$3)))' % (r, r, r, r)).border = border
    ws.cell(row=r, column=12).alignment = center
    ws.cell(row=r, column=13, value='=IF(L%d="","",IF(L%d>=7,"合格","要フォロー"))' % (r, r)).border = border
    ws.cell(row=r, column=13).alignment = center
ws.column_dimensions['A'].width = 16
for i in range(10): ws.column_dimensions[chr(ord('B') + i)].width = 6
ws.column_dimensions['L'].width = 7; ws.column_dimensions['M'].width = 11
setup(ws, landscape=True)

wb.save(out)
print('saved', out)
