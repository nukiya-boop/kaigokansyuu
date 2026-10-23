"""js/data.js から確認テストの解答用紙Excelを作る: python3 tools/make_answer_sheets.py js/data.js docs/確認テスト_解答用紙.xlsx"""
import sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.datavalidation import DataValidation

src, out = sys.argv[1:3]
t = open(src, encoding='utf-8').read()
D = json.loads(t[t.index('{'):t.rindex(';')])

thin = Side(style='thin', color='777777')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill('solid', fgColor='DDEBE6')
input_fill = PatternFill('solid', fgColor='FFFBE6')
wrap = Alignment(wrap_text=True, vertical='center')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)

wb = Workbook()
ws = wb.active; ws.title = '使い方'
lines = [
    D['fiscalYear'] + '　確認テスト 解答用紙（訪問介護）',
    '',
    '・各月のシートが1回分の解答用紙です（A4縦1枚）。印刷して手書きでも、Excelに直接入力しても使えます。',
    '・Excelで入力する場合は、薄い黄色の欄に実施日・氏名を入れ、解答欄は▼から選びます（問1〜5は○×、問6〜10はア・イ・ウ）。',
    '・全10問、1問1点の10点満点です。7点以上が目安です。',
    '・提出された解答は、研修アプリの「紙の解答を採点」に入力すると自動で採点・記録され、解説も表示されます。',
    '・欠席した場合は、研修資料を読んでから解答し、「受験の種類」で「欠席（資料回覧）」に○をつけて提出してください。',
]
for i, l in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=l).alignment = Alignment(wrap_text=True, vertical='top')
ws['A1'].font = Font(bold=True, size=14)
ws.column_dimensions['A'].width = 90
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

for m in D['months']:
    ws = wb.create_sheet(m['month'] + '_解答用紙')
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 64
    ws.column_dimensions['C'].width = 10
    ws.merge_cells('A1:C1')
    ws['A1'] = '確認テスト 解答用紙　' + m['month'] + '　' + m['theme']
    ws['A1'].font = Font(bold=True, size=13); ws['A1'].alignment = wrap
    ws.row_dimensions[1].height = 36
    info = [('実施日', '令和　　年　　月　　日'), ('氏名', ''), ('受験の種類', '研修に参加 ／ 欠席（資料回覧）　※どちらかに○'), ('点数', '　　　点 ／ 10点（採点者記入）')]
    for i, (k, v) in enumerate(info, 3):
        a = ws.cell(row=i, column=1, value=k); a.font = Font(bold=True); a.fill = head_fill; a.border = border; a.alignment = center
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
        b = ws.cell(row=i, column=2, value=v); b.border = border; b.alignment = wrap
        ws.cell(row=i, column=3).border = border
        if k in ('実施日', '氏名'): b.fill = input_fill
        ws.row_dimensions[i].height = 24
    ws.column_dimensions['A'].width = 11
    dv_ox = DataValidation(type='list', formula1='"○,×"', allow_blank=True)
    dv_abc = DataValidation(type='list', formula1='"ア,イ,ウ"', allow_blank=True)
    ws.add_data_validation(dv_ox); ws.add_data_validation(dv_abc)
    r = 8
    for part, rng, note in (('問1〜問5　正しいものに○、誤っているものに×を解答欄に書いてください。', m['test'][:5], dv_ox),
                            ('問6〜問10　正しいもの（または問われているもの）を、ア・イ・ウから1つ選んで解答欄に書いてください。', m['test'][5:], dv_abc)):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        c = ws.cell(row=r, column=1, value=part); c.font = Font(bold=True); c.alignment = wrap
        ws.row_dimensions[r].height = 32
        r += 1
        for i, v in enumerate(['問', '問題', '解答欄'], 1):
            c = ws.cell(row=r, column=i, value=v); c.font = Font(bold=True); c.fill = head_fill; c.border = border; c.alignment = center
        r += 1
        for q in rng:
            text = q['q'] + ''.join('\n　' + ch['k'] + '　' + ch['t'] for ch in q['choices'])
            ws.cell(row=r, column=1, value=q['no']).alignment = center
            ws.cell(row=r, column=2, value=text).alignment = wrap
            a = ws.cell(row=r, column=3); a.fill = input_fill; a.alignment = center; a.font = Font(size=16, bold=True)
            note.add(a.coordinate)
            for col in (1, 2, 3): ws.cell(row=r, column=col).border = border
            nlines = 1 + len(q['choices']) + len(q['q']) // 34
            ws.row_dimensions[r].height = max(30, 17 * nlines)
            r += 1
        r += 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.6, right=0.6, top=0.6, bottom=0.6)
    ws.print_area = 'A1:C%d' % (r - 1)

wb.save(out)
print('saved', out)
