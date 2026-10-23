"""新規採用時研修（虐待防止）の研修資料・確認テストを作る

python3 tools/make_newhire.py js/data.js js/newhire.js docs/新規採用時研修_虐待防止.xlsx

研修資料と確認テストの内容はこのファイルに書いてあります。直したら上のコマンドで
アプリ用のデータ（js/newhire.js）とExcel（研修資料・解答用紙・解答解説）を作り直します。
"""
import sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.datavalidation import DataValidation

src, out_js, out_xlsx = sys.argv[1:4]
t = open(src, encoding='utf-8').read()
D = json.loads(t[t.index('{'):t.rindex(';')])
REFS = [m for m in D['months'] if m['key'] == '05'][0]['refs']

def p(text): return dict(type='p', text=text)
def h(text): return dict(type='h', text=text)
def ox(no, q, a, ex): return dict(no=no, type='ox', q=q, choices=[], answer=a, explain=ex)
def abc(no, q, ch, a, ex): return dict(no=no, type='abc', q=q, choices=[dict(k=k, t=v) for k, v in zip('アイウ', ch)], answer=a, explain=ex)

NH = dict(
    key='nh',
    theme='高齢者虐待防止・身体的拘束等の適正化',
    title='新規採用時研修　高齢者虐待防止・身体的拘束等の適正化（必須）',
    aim='一人で訪問を始める前に、虐待とは何か、してはいけないこと、気づいたときに誰に何を伝えるかを理解する。所要時間の目安は40分（説明25分・確認テスト15分）。',
    blocks=[
        h('1. なぜ入職時に研修を行うのか'),
        p('運営基準により、訪問介護事業所は虐待防止のための研修を、年1回以上に加えて新規採用時にも行うことが義務づけられています。訪問介護は利用者宅で一人で支援することが多く、ほかの職員の目が届きにくいため、一人ひとりが正しい知識を持つことが最も大切な予防策です。'),
        h('2. 高齢者虐待とは'),
        p('高齢者虐待防止法では、65歳以上の高齢者への虐待を「養護者（家族など）による虐待」と「養介護施設従事者等（介護職員など）による虐待」に分けています。訪問介護員は後者にあたります。自分たちが加害者にならないこと、家族などによる虐待に早く気づくことの両方が求められます。'),
        h('3. 虐待の5つの類型'),
        dict(type='table', head=['類型', '内容', '訪問介護で起こりうる例'], rows=[
            ['身体的虐待', '暴力、正当な理由のない身体拘束', '着替えの際に腕を強く引っ張る、つねる'],
            ['介護・世話の放棄・放任', '必要な介護をしない', '排泄介助を時間がないからと省く、おむつを長時間替えない'],
            ['心理的虐待', '暴言、無視、拒絶、恥をかかせる', '「何回言ったらわかるの」と怒鳴る、話しかけを無視する'],
            ['性的虐待', 'わいせつな行為、羞恥心を与える行為', '排泄・入浴介助で不必要に裸のままにする'],
            ['経済的虐待', '財産を不当に使う、使わせない', '買い物のおつりを返さない、利用者の物を持ち帰る'],
        ]),
        h('4. 新しく入った職員が特に気をつけること'),
        p('虐待は「これくらいなら」という不適切ケアの積み重ねから起こります。次のことは行いません。'),
        p('・子ども扱いする呼び方（「〇〇ちゃん」）、命令口調、ため口'),
        p('・「ちょっと待って」と言ったまま待たせる、本人の希望を聞かずに介助を進める'),
        p('・訪問介護計画・手順書にない支援を自分の判断で行う、または省く（迷ったらサービス提供責任者に相談）'),
        p('・利用者宅の物を勝手に使う、持ち帰る。金銭の預かりや買い物の精算を、事業所の決まった手順以外で行う'),
        p('・利用者や家族から金品・お礼を受け取る（事業所のルールに従い、丁寧にお断りする）'),
        p('・同僚の不適切ケアを見ても黙っている'),
        h('5. 気づいたときの報告と通報'),
        p('・虐待かもしれないと思ったら、確証がなくても、その日のうちにサービス提供責任者または虐待防止担当者（管理者）に報告します。'),
        p('・事業所は内容を確認し、市町村へ通報・相談します。事業所に伝えずに、本人が直接市町村へ通報することも認められています。'),
        p('・職員による虐待を発見した場合は、速やかに市町村へ通報する義務があります。家族などによる虐待で、生命・身体に重大な危険が生じている場合も通報義務があり、それ以外は通報の努力義務があります。'),
        p('・通報したことを理由に解雇など不利益な扱いを受けることはありません。守秘義務違反にもなりません。'),
        p('・記録には、見たこと・聞いたことを事実のまま書きます（日時、場所、あざの位置や大きさ、本人や家族の言葉）。推測や評価は書きません。'),
        p('家族などによる虐待のサイン：説明のつかないあざ・傷、急な体重減少、不衛生な状態が続く、本人が家族をおびえる、年金があるのに必要な支払いができない。'),
        h('6. 身体的拘束等は原則禁止'),
        p('令和6年度改定で、訪問介護でも身体的拘束等は原則禁止となり、行う場合は記録が義務づけられました。'),
        p('・例：ベッド柵で四方を囲む、ミトン型手袋、つなぎ服、車いすへのベルト固定、部屋に外から鍵をかける、介助中に両手を押さえ続ける'),
        p('・例外は「切迫性」「非代替性」「一時性」の3要件をすべて満たす緊急やむを得ない場合のみです。自分の判断では行いません。'),
        p('・家族が拘束をしている場面に気づいたら、自己判断で外したり見過ごしたりせず、サービス提供責任者に報告します。'),
        h('7. 当事業所の体制（入職時に記入して説明）'),
        dict(type='table', head=['項目', '内容', '記入欄'], rows=[
            ['虐待防止担当者', '虐待の相談・報告を受ける担当者', ''],
            ['報告先（日中・夜間）', 'サービス提供責任者・管理者の連絡先', ''],
            ['虐待防止のための指針', '保管場所（いつでも読めます）', ''],
            ['虐待防止検討委員会', '定期的に開催し、結果を職員に周知', ''],
            ['職員の相談窓口', '悩みやストレス、ハラスメントの相談先', ''],
        ]),
        h('8. 入職時の確認（理解したら□に✓）'),
        p('□ 虐待の5つの類型と、訪問介護で起こりうる例がわかる'),
        p('□ 不適切ケアをしない、見過ごさないことがわかる'),
        p('□ 気づいたらその日のうちに誰に報告するかがわかる（直接市町村に通報してもよいことも含む）'),
        p('□ 身体的拘束等は原則禁止で、自分の判断では行わないことがわかる'),
        p('□ 虐待防止担当者と、虐待防止のための指針の保管場所がわかる'),
    ],
    test=[
        ox(1, '虐待防止のための研修は、年1回以上に加えて新規採用時にも行う必要がある。', '○', '運営基準で、新規採用時の研修も義務づけられています。'),
        ox(2, '利用者を「〇〇ちゃん」と呼ぶのは親しみの表現なので問題ない。', '×', '子ども扱いの呼び方は尊厳を傷つける不適切ケアで、虐待の入り口になります。'),
        ox(3, '虐待かもしれないと思っても、確証がなければ報告してはいけない。', '×', '確証がなくても、疑いの段階でその日のうちに報告します。'),
        ox(4, '虐待を通報した職員が、それを理由に解雇などの不利益な扱いを受けることはない。', '○', '高齢者虐待防止法で通報者は保護されており、守秘義務違反にもなりません。'),
        ox(5, '身体的拘束等は、家族の同意があればいつでも行ってよい。', '×', '切迫性・非代替性・一時性の3要件をすべて満たす緊急やむを得ない場合のみで、記録も必要です。'),
        abc(6, '訪問介護で起こりうる「身体的虐待」の例はどれか。', ['着替えの際に腕を強く引っ張る', '本人の希望を聞いて服を選ぶ', '介助の前に声をかける'], 'ア', '暴力や無理な力での介助は身体的虐待にあたります。'),
        abc(7, '利用者の腕に説明のつかないあざを見つけたときの対応はどれか。', ['次回の訪問まで様子を見て、誰にも言わない', 'その日のうちにサービス提供責任者（または虐待防止担当者）に報告する', 'その場で家族を問い詰める'], 'イ', '自分だけで判断せず、その日のうちに報告します。家族を問い詰めると、本人がかえって危険になることがあります。'),
        abc(8, '身体的拘束等の例外となる3要件に含まれないものはどれか。', ['切迫性', '非代替性', '経済性'], 'ウ', '3要件は「切迫性」「非代替性」「一時性」です。'),
        abc(9, '利用者から「いつもありがとう」とお礼のお金を渡されたときの対応はどれか。', ['事業所のルールに従い、丁寧にお断りする', '少額なら受け取ってよい', '黙って受け取り、誰にも言わない'], 'ア', '金品の受け取りはトラブルや経済的虐待の入り口になります。断りにくいときはサービス提供責任者に相談します。'),
        abc(10, '気になることがあったときの記録の書き方として適切なものはどれか。', ['「家族が叩いたと思う」と推測を書く', '日時・あざの位置や大きさ・本人の言葉を事実のまま書く', '書かずに口頭で伝えるだけにする'], 'イ', '記録は事実のまま書き、推測や評価は書きません。記録は市町村への通報や対応の根拠になります。'),
    ],
    review=[],
    refs=REFS,
)

with open(out_js, 'w', encoding='utf-8') as f:
    f.write('// 新規採用時研修（虐待防止）の研修資料と確認テスト（tools/make_newhire.py で生成）\n')
    f.write('window.NEWHIRE_ABUSE = ' + json.dumps(NH, ensure_ascii=False, indent=1) + ';\n')

# ---------- Excel ----------
NAVY = '1F3864'
thin = Side(style='thin', color='777777')
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill('solid', fgColor='DDEBE6')
input_fill = PatternFill('solid', fgColor='FFFBE6')
wrap = Alignment(wrap_text=True, vertical='center')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)

def a4(ws, last_col, last_row, one_page):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1 if one_page else 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.6, right=0.6, top=0.6, bottom=0.6)
    ws.print_area = 'A1:%s%d' % (last_col, last_row)

def height(text, width_chars):
    return max(18.75, 16.5 * (1 + len(text) // width_chars))

wb = Workbook()

# 研修資料
ws = wb.active; ws.title = '研修資料'
for k, w in {'A': 24, 'B': 32, 'C': 36}.items(): ws.column_dimensions[k].width = w
r = 1
def line(text, size=10, bold=False, color=None, fill=None):
    global r
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = ws.cell(r, 1, text); c.font = Font(size=size, bold=bold, color=color); c.alignment = wrap
    ws.row_dimensions[r].height = 24 if size >= 14 else height(text, 52)
    r += 1
line(NH['title'], 14, True)
line('ねらい：' + NH['aim'])
for b in NH['blocks']:
    if b['type'] == 'h':
        r += 1; line(b['text'], 11, True, NAVY)
    elif b['type'] == 'p':
        line(b['text'])
    else:
        for i, row in enumerate([b['head']] + b['rows']):
            for j, v in enumerate(row, 1):
                c = ws.cell(r, j, v); c.border = border; c.alignment = wrap; c.font = Font(size=10, bold=(i == 0))
                if i == 0: c.fill = head_fill
                elif b['head'][j - 1] == '記入欄': c.fill = input_fill
            ws.row_dimensions[r].height = max(height(row[k], 18) for k in range(3))
            r += 1
r += 1; line('受講記録（様式5「新規採用時研修記録」に転記）', 11, True, NAVY)
for k, v in [('氏名', ''), ('入職日', '令和　　年　　月　　日'), ('実施日', '令和　　年　　月　　日'), ('講師（職名）', ''),
             ('確認テストの点数', '　　点 ／ 10点'), ('本人署名', ''), ('管理者確認', '')]:
    a = ws.cell(r, 1, k); a.font = Font(size=10, bold=True); a.fill = head_fill; a.border = border; a.alignment = wrap
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    b = ws.cell(r, 2, v); b.border = border; b.alignment = wrap; b.font = Font(size=10)
    ws.cell(r, 3).border = border
    if k not in ('確認テストの点数',): b.fill = input_fill
    ws.row_dimensions[r].height = 24; r += 1
r += 1; line('参考資料（公的機関）', 11, True, NAVY)
for ref in NH['refs']:
    line('・' + ref['name'] + '（' + ref['org'] + '）', color='0563C1')
    ws.cell(r - 1, 1).hyperlink = ref['url']
a4(ws, 'C', r - 1, False)

# 解答用紙
ws = wb.create_sheet('解答用紙')
ws.column_dimensions['A'].width = 11; ws.column_dimensions['B'].width = 64; ws.column_dimensions['C'].width = 10
ws.merge_cells('A1:C1')
ws['A1'] = '確認テスト 解答用紙　新規採用時研修　' + NH['theme']
ws['A1'].font = Font(bold=True, size=13); ws['A1'].alignment = wrap; ws.row_dimensions[1].height = 36
for i, (k, v) in enumerate([('実施日', '令和　　年　　月　　日'), ('氏名', ''), ('入職日', '令和　　年　　月　　日'), ('点数', '　　　点 ／ 10点（採点者記入）')], 3):
    a = ws.cell(row=i, column=1, value=k); a.font = Font(bold=True); a.fill = head_fill; a.border = border; a.alignment = center
    ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=3)
    b = ws.cell(row=i, column=2, value=v); b.border = border; b.alignment = wrap
    ws.cell(row=i, column=3).border = border
    if k != '点数': b.fill = input_fill
    ws.row_dimensions[i].height = 24
dv_ox = DataValidation(type='list', formula1='"○,×"', allow_blank=True)
dv_abc = DataValidation(type='list', formula1='"ア,イ,ウ"', allow_blank=True)
ws.add_data_validation(dv_ox); ws.add_data_validation(dv_abc)
r = 8
for part, rng, dv in (('問1〜問5　正しいものに○、誤っているものに×を解答欄に書いてください。', NH['test'][:5], dv_ox),
                      ('問6〜問10　正しいもの（または問われているもの）を、ア・イ・ウから1つ選んで解答欄に書いてください。', NH['test'][5:], dv_abc)):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    c = ws.cell(row=r, column=1, value=part); c.font = Font(bold=True); c.alignment = wrap; ws.row_dimensions[r].height = 32
    r += 1
    for i, v in enumerate(['問', '問題', '解答欄'], 1):
        c = ws.cell(row=r, column=i, value=v); c.font = Font(bold=True); c.fill = head_fill; c.border = border; c.alignment = center
    r += 1
    for q in rng:
        text = q['q'] + ''.join('\n　' + ch['k'] + '　' + ch['t'] for ch in q['choices'])
        ws.cell(row=r, column=1, value=q['no']).alignment = center
        ws.cell(row=r, column=2, value=text).alignment = wrap
        a = ws.cell(row=r, column=3); a.fill = input_fill; a.alignment = center; a.font = Font(size=16, bold=True)
        dv.add(a.coordinate)
        for col in (1, 2, 3): ws.cell(row=r, column=col).border = border
        ws.row_dimensions[r].height = max(30, 17 * (1 + len(q['choices']) + len(q['q']) // 34))
        r += 1
    r += 1
a4(ws, 'C', r - 1, True)

# 解答・解説
ws = wb.create_sheet('解答・解説')
ws.column_dimensions['A'].width = 6; ws.column_dimensions['B'].width = 8; ws.column_dimensions['C'].width = 70
ws.merge_cells('A1:C1'); ws['A1'] = '確認テスト 解答・解説（管理者・講師用）　新規採用時研修　' + NH['theme']
ws['A1'].font = Font(bold=True, size=13); ws['A1'].alignment = wrap; ws.row_dimensions[1].height = 36
ws.merge_cells('A2:C2'); ws['A2'] = '1問1点の10点満点です。目安は7点以上で、それより低い場合は解説を使って個別に説明します。'
ws['A2'].alignment = wrap
for i, v in enumerate(['問', '正解', '解説'], 1):
    c = ws.cell(row=4, column=i, value=v); c.font = Font(bold=True); c.fill = head_fill; c.border = border; c.alignment = center
r = 5
for q in NH['test']:
    ws.cell(r, 1, q['no']).alignment = center
    c = ws.cell(r, 2, q['answer']); c.alignment = center; c.font = Font(bold=True, size=13)
    ws.cell(r, 3, q['explain']).alignment = wrap
    for col in (1, 2, 3): ws.cell(r, col).border = border
    ws.row_dimensions[r].height = height(q['explain'], 38) + 4
    r += 1
a4(ws, 'C', r - 1, True)

wb.save(out_xlsx)
print('saved', out_js, out_xlsx)
